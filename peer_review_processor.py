#!/usr/bin/env python3
"""
GROUP PEER REVIEW PROCESSOR
===========================

USAGE
-----
This script processes one or more weekly peer-review folders and produces
an Excel report containing group marks, individual marks, attendance,
qualitative feedback, pairings, detected questions, and validation records.

Expected structure for a weekly folder:

    Week_04/
    ├── Group_List.xlsx
    ├── GroupA-GroupB.xlsx
    ├── GroupB-GroupA.xlsx
    ├── GroupC-GroupD.xlsx
    └── ...

A parent directory containing multiple Week_* folders is also supported:

    peer_reviews/
    ├── Week_01/
    ├── Week_02/
    └── Week_03/

Run:

    python peer_review_processor.py ./Week_04

or:

    python peer_review_processor.py ./peer_reviews

Optional custom output path:

    python peer_review_processor.py ./peer_reviews --output report.xlsx

INSTALL DEPENDENCIES
--------------------

    pip install pandas openpyxl

DATA MODEL
----------
1. The roster is the source of truth.
2. Student ID is the primary key.
3. Groups A-H are supported.
4. Each student in a reviewing group submits one review.
5. Each review contains 1-5 quantitative ratings and may contain comments.
6. Different review files/weeks may have different numbers of rating questions.
7. Each reviewer's score is the mean of valid 1-5 rating answers in that file.
8. The reviewed group's mark is the mean of valid reviewer scores received.
9. A rostered student is Present only if they submitted at least one valid review
   for that week. Otherwise they are Absent.
10. Individual Mark /5 = own Group Mark /5 if Present, otherwise 0/5.
11. Qualitative comments never affect marks.
12. No min/max/spread/variance/flags affect the calculation.

ROSTER REQUIREMENTS
-------------------
The roster Excel file must contain columns equivalent to:

    Student ID | Student Name | Group

The script recognises common aliases automatically. The roster filename should
contain one of these terms so it can be detected automatically:

    roster
    group list
    group_list
    grouplist

REVIEW FILE REQUIREMENTS
------------------------
Each review submission must include a Student ID column. Reviewing/reviewed group
may be supplied as spreadsheet columns or inferred from filenames such as:

    GroupA-GroupB.xlsx

meaning:

    Reviewing Group = A
    Reviewed Group  = B

If a submitted reviewing group conflicts with the roster group, the roster wins
and the mismatch is recorded in Validation.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font


# ============================================================================
# CONFIGURATION
# ============================================================================

VALID_GROUPS = set("ABCDEFGH")
MIN_RATING = 1
MAX_RATING = 5

ROSTER_FILENAME_KEYWORDS = (
    "roster",
    "group list",
    "group_list",
    "grouplist",
)

COLUMN_ALIASES = {
    "student_id": [
        "Student ID",
        "Student Id",
        "StudentID",
        "Student Number",
        "Student No",
        "ID",
        "ID Number",
    ],
    "student_name": [
        "Student Name",
        "Name",
        "Full Name",
        "Reviewer Name",
        "Your Name",
    ],
    "group": [
        "Group",
        "Student Group",
        "Your Group",
        "Group Letter",
    ],
    "reviewing_group": [
        "Reviewing Group",
        "Reviewer Group",
        "Your Group",
        "Your Group Letter",
        "Group",
    ],
    "reviewed_group": [
        "Reviewed Group",
        "Group Reviewed",
        "Group Being Reviewed",
        "Which group are you reviewing?",
        "Peer Group",
        "Review Group",
        "Group to Review",
    ],
}

ADMIN_KEYWORDS = [
    "timestamp",
    "start time",
    "completion time",
    "email",
    "e-mail",
    "student id",
    "student number",
    "student no",
    "id number",
    "username",
    "user name",
    "first name",
    "last name",
    "full name",
    "student name",
    "reviewer name",
    "reviewing group",
    "reviewed group",
    "group being reviewed",
    "which group",
    "source file",
    "week",
]

QUALITATIVE_KEYWORDS = [
    "comment",
    "feedback",
    "explain",
    "description",
    "describe",
    "strength",
    "improve",
    "improvement",
    "suggest",
    "suggestion",
    "reason",
    "observation",
    "recommend",
    "recommendation",
    "what worked",
    "what could",
    "what went well",
    "additional",
    "elaborate",
    "justify",
]


# ============================================================================
# TEXT / COLUMN HELPERS
# ============================================================================


def normalise_text(value: object) -> str:
    """Return lowercase, whitespace-normalised text."""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def find_column(df: pd.DataFrame, aliases: Iterable[str]) -> str | None:
    """Find a dataframe column matching one of the supplied aliases."""
    lookup = {normalise_text(col): col for col in df.columns}
    for alias in aliases:
        key = normalise_text(alias)
        if key in lookup:
            return lookup[key]
    return None


def clean_student_id(value: object) -> str | None:
    """Normalise Student ID to a stable string primary key."""
    if pd.isna(value):
        return None

    # Excel often turns numeric IDs into floats (e.g. 12345678.0).
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)) and float(value).is_integer():
        return str(int(value))

    text = str(value).strip()
    if not text:
        return None

    # Remove a trailing .0 produced by spreadsheet imports.
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]

    return text


def extract_group_letter(value: object) -> str | None:
    """Extract and validate a group letter A-H."""
    if pd.isna(value):
        return None

    text = str(value).strip().upper()
    if text in VALID_GROUPS:
        return text

    match = re.search(r"\bGROUP[\s_-]*([A-H])\b", text, re.IGNORECASE)
    if match:
        return match.group(1).upper()

    return None


def group_label(letter: str | None) -> str | None:
    return f"Group {letter}" if letter in VALID_GROUPS else None


def looks_administrative(column_name: object) -> bool:
    name = normalise_text(column_name)
    return any(keyword in name for keyword in ADMIN_KEYWORDS)


def looks_qualitative(column_name: object) -> bool:
    name = normalise_text(column_name)
    return any(keyword in name for keyword in QUALITATIVE_KEYWORDS)


# ============================================================================
# WEEK / FILE DETECTION
# ============================================================================


def extract_week_number(text: str) -> int | None:
    match = re.search(r"week[\s_-]*(\d+)", text, re.IGNORECASE)
    return int(match.group(1)) if match else None


def detect_week(folder: Path) -> int | str:
    """Detect week from folder name; otherwise keep folder name as identifier."""
    week = extract_week_number(folder.name)
    return week if week is not None else folder.name


def is_roster_file(path: Path) -> bool:
    stem = normalise_text(path.stem).replace("-", " ")
    return any(keyword in stem for keyword in ROSTER_FILENAME_KEYWORDS)


def find_roster_file(folder: Path) -> Path:
    candidates = [
        p for p in folder.glob("*.xlsx")
        if not p.name.startswith("~$") and is_roster_file(p)
    ]

    if len(candidates) == 1:
        return candidates[0]
    if not candidates:
        raise FileNotFoundError(
            f"No roster/group-list Excel file found in: {folder}\n"
            "Rename the roster so its filename contains 'roster' or 'group list'."
        )

    names = ", ".join(p.name for p in candidates)
    raise ValueError(f"Multiple roster files found in {folder}: {names}")


def detect_pairing_from_filename(path: Path) -> tuple[str | None, str | None]:
    """
    Detect filenames such as:
        GroupA-GroupB.xlsx
        Group A - Group B.xlsx
        A-B.xlsx
    Returns (reviewing_group_letter, reviewed_group_letter).
    """
    stem = path.stem.upper()

    patterns = [
        r"GROUP[\s_-]*([A-H]).*?GROUP[\s_-]*([A-H])",
        r"^\s*([A-H])\s*[-_]+\s*([A-H])\s*$",
    ]

    for pattern in patterns:
        match = re.search(pattern, stem, re.IGNORECASE)
        if match:
            return match.group(1).upper(), match.group(2).upper()

    return None, None


def discover_week_folders(input_path: Path) -> list[Path]:
    """Support either a single week folder or a parent of Week_* folders."""
    if not input_path.exists() or not input_path.is_dir():
        raise FileNotFoundError(f"Input folder does not exist: {input_path}")

    # If this folder already contains a roster, treat it as one week.
    if any(is_roster_file(p) for p in input_path.glob("*.xlsx")):
        return [input_path]

    week_folders = sorted(
        [p for p in input_path.iterdir() if p.is_dir() and extract_week_number(p.name) is not None],
        key=lambda p: extract_week_number(p.name) or 0,
    )

    if not week_folders:
        raise FileNotFoundError(
            f"No weekly folders found under {input_path}. "
            "Expected Week_01, Week_02, ... or a single folder containing a roster."
        )

    return week_folders


# ============================================================================
# QUESTION DETECTION
# ============================================================================


def identify_rating_columns(df: pd.DataFrame, excluded: set[str]) -> list[str]:
    """
    Detect quantitative rating columns independently for each review file.
    A rating column must be mostly numeric and every numeric value must be 1-5.
    """
    rating_columns: list[str] = []

    for col in df.columns:
        if col in excluded or looks_administrative(col) or looks_qualitative(col):
            continue

        original_non_empty = df[col].replace("", np.nan).dropna()
        if original_non_empty.empty:
            continue

        numeric = pd.to_numeric(df[col], errors="coerce")
        valid = numeric.dropna()
        if valid.empty:
            continue

        numeric_ratio = len(valid) / len(original_non_empty)
        if numeric_ratio < 0.80:
            continue

        if valid.between(MIN_RATING, MAX_RATING, inclusive="both").all():
            rating_columns.append(col)

    return rating_columns


def identify_qualitative_columns(
    df: pd.DataFrame,
    excluded: set[str],
    rating_columns: list[str],
) -> list[str]:
    qualitative: list[str] = []
    for col in df.columns:
        if col in excluded or col in rating_columns:
            continue
        if looks_qualitative(col):
            qualitative.append(col)
    return qualitative


# ============================================================================
# ROSTER
# ============================================================================


def load_roster(roster_path: Path, week: int | str) -> tuple[pd.DataFrame, list[dict]]:
    """Load, validate, and standardise the weekly roster."""
    df = pd.read_excel(roster_path).dropna(how="all").copy()

    id_col = find_column(df, COLUMN_ALIASES["student_id"])
    name_col = find_column(df, COLUMN_ALIASES["student_name"])
    group_col = find_column(df, COLUMN_ALIASES["group"])

    if id_col is None or name_col is None or group_col is None:
        raise ValueError(
            f"Roster {roster_path.name} must contain Student ID, Student Name, and Group columns."
        )

    roster = pd.DataFrame({
        "Week": week,
        "Student ID": df[id_col].apply(clean_student_id),
        "Student Name": df[name_col].fillna("").astype(str).str.strip(),
        "Group": df[group_col].apply(extract_group_letter),
    })

    validation: list[dict] = []

    for idx, row in roster.iterrows():
        if not row["Student ID"]:
            validation.append({
                "Week": week,
                "Source File": roster_path.name,
                "Student ID": "",
                "Issue": "Roster row missing Student ID",
                "Details": f"Row {idx + 2}",
            })
        if row["Group"] not in VALID_GROUPS:
            validation.append({
                "Week": week,
                "Source File": roster_path.name,
                "Student ID": row["Student ID"] or "",
                "Issue": "Invalid roster group",
                "Details": f"Expected Group A-H; row {idx + 2}",
            })

    roster = roster[roster["Student ID"].notna() & roster["Group"].isin(VALID_GROUPS)].copy()

    duplicate_ids = roster[roster.duplicated("Student ID", keep=False)]
    if not duplicate_ids.empty:
        dupes = ", ".join(sorted(duplicate_ids["Student ID"].unique()))
        raise ValueError(f"Duplicate Student ID(s) in roster {roster_path.name}: {dupes}")

    return roster.reset_index(drop=True), validation


# ============================================================================
# REVIEW PROCESSING
# ============================================================================


def process_review_file(
    review_path: Path,
    roster: pd.DataFrame,
    week: int | str,
) -> tuple[pd.DataFrame, list[dict], list[dict], list[dict]]:
    """Process one review workbook and return reviews, comments, detection, validation."""
    df = pd.read_excel(review_path).dropna(how="all").copy()

    if df.empty:
        return pd.DataFrame(), [], [], [{
            "Week": week,
            "Source File": review_path.name,
            "Student ID": "",
            "Issue": "Empty review file",
            "Details": "No response rows found",
        }]

    student_id_col = find_column(df, COLUMN_ALIASES["student_id"])
    reviewing_group_col = find_column(df, COLUMN_ALIASES["reviewing_group"])
    reviewed_group_col = find_column(df, COLUMN_ALIASES["reviewed_group"])

    if student_id_col is None:
        raise ValueError(f"No Student ID column found in {review_path.name}")

    file_reviewing, file_reviewed = detect_pairing_from_filename(review_path)

    if reviewed_group_col is None and file_reviewed is None:
        raise ValueError(
            f"Could not identify reviewed group in {review_path.name}. "
            "Add a Reviewed Group column or use a filename such as GroupA-GroupB.xlsx."
        )

    excluded = {student_id_col}
    for col in (reviewing_group_col, reviewed_group_col):
        if col:
            excluded.add(col)

    rating_columns = identify_rating_columns(df, excluded)
    qualitative_columns = identify_qualitative_columns(df, excluded, rating_columns)

    detection_rows: list[dict] = []
    for i, col in enumerate(rating_columns, 1):
        detection_rows.append({
            "Week": week,
            "Source File": review_path.name,
            "Question Type": "Rating",
            "Question Number": i,
            "Question": col,
        })
    for i, col in enumerate(qualitative_columns, 1):
        detection_rows.append({
            "Week": week,
            "Source File": review_path.name,
            "Question Type": "Qualitative",
            "Question Number": i,
            "Question": col,
        })

    validation: list[dict] = []
    comments: list[dict] = []
    processed_rows: list[dict] = []

    roster_lookup = roster.set_index("Student ID").to_dict("index")

    for excel_row_index, row in df.iterrows():
        student_id = clean_student_id(row.get(student_id_col))

        submitted_reviewing = (
            extract_group_letter(row.get(reviewing_group_col)) if reviewing_group_col else file_reviewing
        )
        submitted_reviewed = (
            extract_group_letter(row.get(reviewed_group_col)) if reviewed_group_col else file_reviewed
        )

        if student_id is None:
            validation.append({
                "Week": week,
                "Source File": review_path.name,
                "Student ID": "",
                "Issue": "Missing Student ID",
                "Details": f"Excel row {excel_row_index + 2}",
            })
            continue

        if student_id not in roster_lookup:
            validation.append({
                "Week": week,
                "Source File": review_path.name,
                "Student ID": student_id,
                "Issue": "Unknown Student ID",
                "Details": "Student ID not found in weekly roster",
            })
            continue

        roster_record = roster_lookup[student_id]
        roster_group = roster_record["Group"]
        student_name = roster_record["Student Name"]

        # Roster is authoritative for reviewing group.
        reviewing_group = roster_group

        if submitted_reviewing and submitted_reviewing != roster_group:
            validation.append({
                "Week": week,
                "Source File": review_path.name,
                "Student ID": student_id,
                "Issue": "Submitted group differs from roster",
                "Details": f"Submitted={submitted_reviewing}; Roster={roster_group}; roster used",
            })

        reviewed_group = submitted_reviewed
        if reviewed_group not in VALID_GROUPS:
            validation.append({
                "Week": week,
                "Source File": review_path.name,
                "Student ID": student_id,
                "Issue": "Invalid or missing reviewed group",
                "Details": f"Excel row {excel_row_index + 2}",
            })
            continue

        if reviewed_group == reviewing_group:
            validation.append({
                "Week": week,
                "Source File": review_path.name,
                "Student ID": student_id,
                "Issue": "Student reviewing own group",
                "Details": f"Group {reviewing_group} -> Group {reviewed_group}",
            })
            continue

        rating_values = []
        for col in rating_columns:
            value = pd.to_numeric(pd.Series([row.get(col)]), errors="coerce").iloc[0]
            if pd.notna(value) and MIN_RATING <= float(value) <= MAX_RATING:
                rating_values.append(float(value))

        reviewer_score = round(float(np.mean(rating_values)), 2) if rating_values else np.nan

        if not rating_values:
            validation.append({
                "Week": week,
                "Source File": review_path.name,
                "Student ID": student_id,
                "Issue": "No valid quantitative rating",
                "Details": "Submission does not count as attendance or mark evidence",
            })

        processed_rows.append({
            "Week": week,
            "Student ID": student_id,
            "Student Name": student_name,
            "Reviewing Group": group_label(reviewing_group),
            "Reviewed Group": group_label(reviewed_group),
            "Rating Questions Detected": len(rating_columns),
            "Rating Questions Answered": len(rating_values),
            "Reviewer Score /5": reviewer_score,
            "Valid Submission": bool(rating_values),
            "Source File": review_path.name,
            "Source Row": excel_row_index + 2,
        })

        for question in qualitative_columns:
            response = row.get(question)
            if pd.isna(response) or not str(response).strip():
                continue
            comments.append({
                "Week": week,
                "Student ID": student_id,
                "Student Name": student_name,
                "Reviewing Group": group_label(reviewing_group),
                "Reviewed Group": group_label(reviewed_group),
                "Question": question,
                "Response": str(response).strip(),
                "Source File": review_path.name,
            })

    return pd.DataFrame(processed_rows), comments, detection_rows, validation


# ============================================================================
# WEEK PROCESSING
# ============================================================================


def process_week_folder(folder: Path):
    week = detect_week(folder)
    roster_path = find_roster_file(folder)
    roster, validation = load_roster(roster_path, week)

    review_files = [
        p for p in folder.glob("*.xlsx")
        if not p.name.startswith("~$") and p != roster_path
    ]

    if not review_files:
        raise FileNotFoundError(f"No review Excel files found in {folder}")

    reviews_frames: list[pd.DataFrame] = []
    comments: list[dict] = []
    detection: list[dict] = []

    for review_path in sorted(review_files):
        reviews, file_comments, file_detection, file_validation = process_review_file(
            review_path, roster, week
        )
        if not reviews.empty:
            reviews_frames.append(reviews)
        comments.extend(file_comments)
        detection.extend(file_detection)
        validation.extend(file_validation)

    reviews = pd.concat(reviews_frames, ignore_index=True) if reviews_frames else pd.DataFrame()

    # Duplicate valid submissions from the same student for the same pairing.
    if not reviews.empty:
        duplicate_mask = reviews.duplicated(
            subset=["Week", "Student ID", "Reviewing Group", "Reviewed Group"],
            keep=False,
        ) & reviews["Valid Submission"]

        duplicates = reviews[duplicate_mask]
        for _, row in duplicates.iterrows():
            validation.append({
                "Week": week,
                "Source File": row["Source File"],
                "Student ID": row["Student ID"],
                "Issue": "Duplicate valid submission",
                "Details": f"{row['Reviewing Group']} -> {row['Reviewed Group']}",
            })

        # For calculation, keep only the last valid submission per student/pairing.
        # Invalid submissions do not override a valid one.
        valid_reviews = reviews[reviews["Valid Submission"]].copy()
        valid_reviews = valid_reviews.drop_duplicates(
            subset=["Week", "Student ID", "Reviewing Group", "Reviewed Group"],
            keep="last",
        )
    else:
        valid_reviews = pd.DataFrame()

    # Pairings are based on valid submissions only.
    if not valid_reviews.empty:
        pairings = (
            valid_reviews[["Week", "Reviewing Group", "Reviewed Group"]]
            .drop_duplicates()
            .sort_values(["Week", "Reviewing Group", "Reviewed Group"])
            .reset_index(drop=True)
        )
    else:
        pairings = pd.DataFrame(columns=["Week", "Reviewing Group", "Reviewed Group"])

    # Group marks.
    if not valid_reviews.empty:
        group_marks = (
            valid_reviews.groupby(
                ["Week", "Reviewing Group", "Reviewed Group"], dropna=False
            )
            .agg(
                **{
                    "Reviews Received": ("Reviewer Score /5", "count"),
                    "Mean /5": ("Reviewer Score /5", "mean"),
                }
            )
            .reset_index()
        )
        group_marks["Mean /5"] = group_marks["Mean /5"].round(2)

        # Expected reviews = roster size of reviewing group.
        group_sizes = roster.groupby("Group")["Student ID"].nunique().to_dict()
        group_marks["Reviews Expected"] = group_marks["Reviewing Group"].apply(
            lambda label: group_sizes.get(extract_group_letter(label), 0)
        )
        group_marks = group_marks[
            [
                "Week",
                "Reviewing Group",
                "Reviewed Group",
                "Reviews Expected",
                "Reviews Received",
                "Mean /5",
            ]
        ]
    else:
        group_marks = pd.DataFrame(columns=[
            "Week", "Reviewing Group", "Reviewed Group", "Reviews Expected", "Reviews Received", "Mean /5"
        ])

    # Map each group's received mark: Reviewed Group -> Mean /5.
    own_group_mark_map: dict[str, float] = {}
    for _, row in group_marks.iterrows():
        reviewed_letter = extract_group_letter(row["Reviewed Group"])
        if reviewed_letter:
            # If multiple pairings review the same group, average them rather than overwrite.
            own_group_mark_map.setdefault(reviewed_letter, [])
            own_group_mark_map[reviewed_letter].append(float(row["Mean /5"]))
    own_group_mark_map = {
        group: round(float(np.mean(values)), 2)
        for group, values in own_group_mark_map.items()
    }

    # Attendance = at least one valid review submitted that week.
    present_ids = set(valid_reviews["Student ID"].astype(str)) if not valid_reviews.empty else set()

    individual_rows = []
    attendance_rows = []
    for _, student in roster.iterrows():
        sid = student["Student ID"]
        group = student["Group"]
        present = sid in present_ids
        group_mark = own_group_mark_map.get(group, np.nan)
        individual_mark = round(float(group_mark), 2) if present and pd.notna(group_mark) else 0.0

        row = {
            "Week": week,
            "Student ID": sid,
            "Student Name": student["Student Name"],
            "Group": group_label(group),
            "Attendance": "Present" if present else "Absent",
            "Group Mark /5": group_mark,
            "Individual Mark /5": individual_mark,
        }
        individual_rows.append(row)
        attendance_rows.append({
            "Week": week,
            "Student ID": sid,
            "Student Name": student["Student Name"],
            "Group": group_label(group),
            "Attendance": "Present" if present else "Absent",
        })

    individual_marks = pd.DataFrame(individual_rows)
    attendance = pd.DataFrame(attendance_rows)

    return {
        "roster": roster,
        "group_marks": group_marks,
        "individual_marks": individual_marks,
        "attendance": attendance,
        "individual_reviews": reviews,
        "pairings": pairings,
        "comments": pd.DataFrame(comments),
        "question_detection": pd.DataFrame(detection),
        "validation": pd.DataFrame(validation),
    }


# ============================================================================
# OUTPUT
# ============================================================================


def concat_or_empty(frames: list[pd.DataFrame], columns: list[str] | None = None) -> pd.DataFrame:
    usable = [df for df in frames if df is not None and not df.empty]
    if usable:
        return pd.concat(usable, ignore_index=True, sort=False)
    return pd.DataFrame(columns=columns or [])


def auto_format_workbook(output_path: Path) -> None:
    workbook = load_workbook(output_path)

    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top")

        for column_cells in ws.columns:
            column_letter = column_cells[0].column_letter
            max_length = 0
            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = min(max_length + 2, 70)

        # Wrap all cells on text-heavy sheets.
        if ws.title in {"Qualitative Feedback", "Validation", "Question Detection"}:
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.save(output_path)


def write_report(results: list[dict], output_path: Path) -> None:
    group_marks = concat_or_empty([r["group_marks"] for r in results])
    individual_marks = concat_or_empty([r["individual_marks"] for r in results])
    attendance = concat_or_empty([r["attendance"] for r in results])
    individual_reviews = concat_or_empty([r["individual_reviews"] for r in results])
    pairings = concat_or_empty([r["pairings"] for r in results])
    comments = concat_or_empty([r["comments"] for r in results])
    question_detection = concat_or_empty([r["question_detection"] for r in results])
    validation = concat_or_empty([r["validation"] for r in results])

    # Sort primary reports for readability.
    if not group_marks.empty:
        group_marks = group_marks.sort_values(["Week", "Reviewed Group", "Reviewing Group"])
    if not individual_marks.empty:
        individual_marks = individual_marks.sort_values(["Week", "Group", "Student Name", "Student ID"])
    if not attendance.empty:
        attendance = attendance.sort_values(["Week", "Group", "Student Name", "Student ID"])
    if not individual_reviews.empty:
        individual_reviews = individual_reviews.sort_values(["Week", "Reviewed Group", "Student ID"])

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        group_marks.to_excel(writer, sheet_name="Group Marks", index=False)
        individual_marks.to_excel(writer, sheet_name="Individual Marks", index=False)
        attendance.to_excel(writer, sheet_name="Attendance", index=False)
        individual_reviews.to_excel(writer, sheet_name="Individual Reviews", index=False)
        comments.to_excel(writer, sheet_name="Qualitative Feedback", index=False)
        pairings.to_excel(writer, sheet_name="Pairings", index=False)
        question_detection.to_excel(writer, sheet_name="Question Detection", index=False)
        validation.to_excel(writer, sheet_name="Validation", index=False)

    auto_format_workbook(output_path)


# ============================================================================
# CLI
# ============================================================================


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Process weekly group peer-review Excel files and generate group/individual marks."
    )
    parser.add_argument(
        "input_folder",
        type=Path,
        help="Single week folder or parent folder containing Week_* folders",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("peer_review_report.xlsx"),
        help="Output Excel workbook path (default: peer_review_report.xlsx)",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()

    try:
        week_folders = discover_week_folders(args.input_folder)
        results = []

        for folder in week_folders:
            print(f"Processing {folder} ...")
            results.append(process_week_folder(folder))

        write_report(results, args.output)

        total_group_rows = sum(len(r["group_marks"]) for r in results)
        total_students = sum(len(r["individual_marks"]) for r in results)
        total_validation = sum(len(r["validation"]) for r in results)

        print("\nProcessing complete")
        print(f"Weeks processed: {len(results)}")
        print(f"Group-mark rows: {total_group_rows}")
        print(f"Individual-mark rows: {total_students}")
        print(f"Validation records: {total_validation}")
        print(f"Report: {args.output.resolve()}")
        return 0

    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
